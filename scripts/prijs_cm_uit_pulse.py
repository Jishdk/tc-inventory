"""prijs_cm (kolom M) in `Inventory TC.xlsx` bijwerken vanuit de Pulse-adviesprijzen,
op XML-niveau zodat formules, caches en opmaak intact blijven (nooit openpyxl-save).

    python scripts/prijs_cm_uit_pulse.py ADVIES.json [--apply]
        [--bestand PAD] [--skip NAAM[;CODE]]... [--comp-override NAAM[;CODE]=BEDRAG]...
        [--prijs NAAM[;CODE]=BEDRAG]...

ADVIES.json is een lijst van {naam, code, cat, prijs_cm, advies} zoals de Pulse-ronde
die oplevert (`naam`/`code`/`cat` = Excel-kolommen A/B/E, `prijs_cm` = de huidige
M-waarde als controle, `advies` = de nieuwe prijs). Rijen worden gezocht op
naam+code+categorie — niet op `items.bron_rij`, want de Windows-download schuift
t.o.v. de database (16-10-2026: Ho-Oh Ex 230/182 op rij 936 vs bron_rij 941).

Uitzonderingen (de drie regels van de eerste ronde, 16-10-2026):
  --skip NAAM            rij overslaan (blijft op de oude prijs)
  --comp-override N=B    kolom N wordt een vaste comp B: de formule gaat weg (ook uit
                         calcChain). Nodig boven €700, daar geeft de ladder "".
  --prijs N=B            nieuwe prijs B in plaats van het advies; is NAAM niet in het
                         advies-bestand, dan wordt de rij toegevoegd (code/cat worden
                         dan uit de Excel gelezen; de naam moet dan uniek zijn).
NAAM matcht hoofdletterongevoelig op kolom A; is dat niet uniek binnen het advies,
geef dan NAAM;CODE.

Zonder --apply alleen de tabel oud -> nieuw. Met --apply: back-up in data/backup/,
M bijwerken, caches N (comp-ladder, alleen waar N een formule is), O en P en het
Dashboard herrekenen, fullCalcOnLoad=1, integriteitscontrole (alle andere cellen
identiek), dan pas het origineel vervangen.
"""
from __future__ import annotations

import argparse
import datetime
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from superclean_v7 import CEL_RE, REF_RE, comp_ladder, zet_cache  # noqa: E402
from weekend_v7_rosmalen import (  # noqa: E402
    DASH, SHEET, controleer, herreken_dashboard, kolom_data, lees, schrijf,
)

INV = Path(__file__).resolve().parent.parent
STANDAARD_BESTAND = Path("/mnt/c/Users/jishn/Documents/Projects/TraderCollective/Inventory TC.xlsx")
BACKUP_DIR = INV / "data" / "backup"
M, N, O, P = "M", "N", "O", "P"


def num(v):
    return None if v in (None, "") else v


def sleutel(naam, code=None):
    return (str(naam or "").strip().lower(), None if code is None else str(code).strip().lower())


def parse_naam(spec: str):
    """'Gyarados' -> ('gyarados', None); 'Pikachu van Gogh;085' -> ('pikachu van gogh', '085')."""
    naam, _, code = spec.partition(";")
    return sleutel(naam, code or None)


def parse_bedrag(spec: str):
    links, _, bedrag = spec.rpartition("=")
    if not links or not bedrag:
        sys.exit(f"verwacht NAAM[;CODE]=BEDRAG, kreeg {spec!r}")
    return parse_naam(links), float(bedrag.replace(",", "."))


def past(o: dict, key) -> bool:
    naam, code = key
    return sleutel(o["naam"])[0] == naam and (code is None or sleutel(None, o.get("code"))[1] == code)


def selecteer(cands: list[dict], key, wat: str) -> list[dict]:
    hits = [o for o in cands if past(o, key)]
    if len(hits) > 1:
        sys.exit(f"{wat}: {key[0]!r} matcht {len(hits)} rijen in het advies — geef NAAM;CODE")
    return hits


def cel(cellen: dict, kol: str) -> str:
    return next(c for ref, c in cellen.items() if ref.rstrip("0123456789") == kol)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("advies", type=Path)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--bestand", type=Path, default=STANDAARD_BESTAND)
    ap.add_argument("--skip", action="append", default=[], metavar="NAAM[;CODE]")
    ap.add_argument("--comp-override", action="append", default=[], metavar="NAAM[;CODE]=BEDRAG")
    ap.add_argument("--prijs", action="append", default=[], metavar="NAAM[;CODE]=BEDRAG")
    args = ap.parse_args()

    cands = json.load(open(args.advies))
    wb = openpyxl.load_workbook(args.bestand, data_only=True)
    inv = wb["Inventaris"]
    index: dict[tuple, list[int]] = {}
    for r in range(2, inv.max_row + 1):
        a, b, e = inv.cell(r, 1).value, inv.cell(r, 2).value, inv.cell(r, 5).value
        if a:
            index.setdefault((*sleutel(a, b if b is not None else ""), str(e or "").strip().lower()), []).append(r)

    # -- uitzonderingen -------------------------------------------------------
    for spec in args.skip:
        key = parse_naam(spec)
        hits = selecteer(cands, key, "--skip")
        if not hits:
            sys.exit(f"--skip: {spec!r} staat niet in het advies")
        cands = [o for o in cands if o is not hits[0]]
        print(f"overgeslagen: {hits[0]['naam']} {hits[0].get('code') or ''} (blijft {hits[0]['prijs_cm']:.0f})")
    for spec in args.prijs:
        key, bedrag = parse_bedrag(spec)
        hits = selecteer(cands, key, "--prijs")
        if hits:
            hits[0]["advies"] = bedrag
        else:  # rij toevoegen vanuit de Excel
            rijen = [r for (a, b, e), rs in index.items() if a == key[0] and (key[1] is None or b == key[1]) for r in rs]
            if len(rijen) != 1:
                sys.exit(f"--prijs: {spec!r} niet in het advies en {len(rijen)} rijen in de Excel ({rijen})")
            r = rijen[0]
            cands.append(dict(naam=inv.cell(r, 1).value, code=inv.cell(r, 2).value, cat=inv.cell(r, 5).value,
                              prijs_cm=float(inv.cell(r, 13).value), advies=bedrag))
    for spec in args.comp_override:
        key, bedrag = parse_bedrag(spec)
        hits = selecteer(cands, key, "--comp-override")
        if not hits:
            sys.exit(f"--comp-override: {spec!r} staat niet in het advies")
        hits[0]["comp_vast"] = bedrag

    # -- plan: elke adviesrij naar precies één Excel-rij ------------------------
    plan, by_row = [], {}
    for o in cands:
        key = (*sleutel(o["naam"], o.get("code") or ""), str(o["cat"] or "").strip().lower())
        rows = index.get(key, [])
        if len(rows) != 1:
            sys.exit(f"{o['naam']} {o.get('code')}: {len(rows)} rijen in Excel ({rows}) — stop")
        r = rows[0]
        oud = inv.cell(r, 13).value
        if num(oud) is None or float(oud) != float(o["prijs_cm"]):
            sys.exit(f"rij {r} {o['naam']}: M={oud!r} maar advies verwacht prijs_cm={o['prijs_cm']} — stop")
        by_row[r] = o
        plan.append(dict(rij=r, naam=inv.cell(r, 1).value, code=inv.cell(r, 2).value, cat=inv.cell(r, 5).value,
                         j=inv.cell(r, 10).value, oud=float(oud), nieuw=float(o["advies"]), n_oud=inv.cell(r, 14).value))

    # -- XML -----------------------------------------------------------------
    print(f"{'rij':>4} {'naam':28} {'code':10} {'cat':6} {'oud':>8} {'nieuw':>6} {'comp oud':>9} {'comp nieuw':>10}")
    kol = kolom_data(inv)
    stukken = lees(args.bestand)
    xml = stukken[SHEET].decode("utf-8")
    wijzigingen, calc_weg = [], []
    for p in plan:
        rn, o = p["rij"], by_row[p["rij"]]
        rij_xml = re.search(rf'<row r="{rn}"[^>]*>.*?</row>', xml, re.S).group(0)
        cellen = {REF_RE.search(c.group(0)).group(1): c.group(0) for c in CEL_RE.finditer(rij_xml)}
        m_cel, n_cel, o_cel, p_cel = (cel(cellen, k) for k in (M, N, O, P))
        if "<f" in m_cel:
            sys.exit(f"rij {rn}: M is een formule — stop")
        nieuw_getal = f"{p['nieuw']:g}"
        nieuw_m = re.sub(r"<v>[^<]*</v>", f"<v>{nieuw_getal}</v>", m_cel)
        j = num(p["j"]) or 0
        if o.get("comp_vast") is not None:
            if 't="shared"' in n_cel:
                sys.exit(f"rij {rn}: N is een gedeelde formule (master?) — handmatig")
            attrs = re.sub(r'\st="[^"]*"', "", re.match(r"<c\b([^>]*?)/?>", n_cel).group(1)).rstrip()
            comp = float(o["comp_vast"])
            nieuw_n = f"<c{attrs}><v>{comp:g}</v></c>"
            p["comp_nieuw"] = f"{comp:g} (vast, formule weg)"
            if "<f" in n_cel:
                calc_weg.append(f"N{rn}")
        elif "<f" in n_cel:
            comp = comp_ladder(p["nieuw"])
            nieuw_n = zet_cache(n_cel, comp)
            p["comp_nieuw"] = f"{comp:g}" if comp != "" else "(leeg: ladder stopt bij 700)"
        else:
            comp, nieuw_n = num(p["n_oud"]), n_cel
            p["comp_nieuw"] = f"{comp} (vast)"
        nieuw_o = zet_cache(o_cel, float(j * p["nieuw"])) if "<f" in o_cel else o_cel
        p_val = float(j * comp) if comp not in ("", None) else float(j * p["nieuw"])
        nieuw_p = zet_cache(p_cel, p_val) if "<f" in p_cel else p_cel
        nieuw_rij = (rij_xml.replace(m_cel, nieuw_m).replace(n_cel, nieuw_n)
                     .replace(o_cel, nieuw_o).replace(p_cel, nieuw_p))
        controleer(nieuw_rij, list(cellen), rn)
        wijzigingen.append((rij_xml, nieuw_rij))
        kol[M][rn - 1] = p["nieuw"]
        kol[N][rn - 1] = comp if comp not in ("", None) else None
        print(f"{rn:4} {str(p['naam'])[:28]:28} {str(p['code'] or ''):10} {str(p['cat']):6} "
              f"{p['oud']:8.0f} {p['nieuw']:6g} {str(p['n_oud']):>9} {p['comp_nieuw']:>10}")
    if not args.apply:
        print(f"\n(dry-run: {len(plan)} rijen, niets geschreven)")
        return

    for oud, nieuw in wijzigingen:
        assert xml.count(oud) == 1
        xml = xml.replace(oud, nieuw)
    stukken[SHEET] = xml.encode("utf-8")
    if calc_weg:
        cc = stukken["xl/calcChain.xml"].decode("utf-8")
        for ref in calc_weg:
            n_voor = cc.count("<c ")
            cc = re.sub(rf'<c [^>]*r="{ref}"[^>]*/>', "", cc)
            assert cc.count("<c ") == n_voor - 1, ref
        stukken["xl/calcChain.xml"] = cc.encode("utf-8")
    herreken_dashboard(stukken, kol)
    wbx = stukken["xl/workbook.xml"].decode("utf-8")
    stukken["xl/workbook.xml"] = re.sub(r"<calcPr[^>]*/>", '<calcPr fullCalcOnLoad="1"/>', wbx).encode("utf-8")

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    backup = BACKUP_DIR / f"InventoryTC_{stamp}_voor_pulse_prijs_cm.xlsx"
    shutil.copy2(args.bestand, backup)
    tmp = BACKUP_DIR / f"InventoryTC_{stamp}_na_pulse_prijs_cm.tmp.xlsx"
    schrijf(stukken, tmp)
    # integriteit: leesbaar, alle rijen, en buiten M/N/O/P van de planrijen is niets veranderd
    chk = openpyxl.load_workbook(tmp, data_only=True)["Inventaris"]
    assert chk.max_row == inv.max_row, (chk.max_row, inv.max_row)
    planrijen = {p["rij"] for p in plan}
    for r in range(2, inv.max_row + 1):
        for c in range(1, 20):
            if r in planrijen and c in (13, 14, 15, 16):
                continue
            assert chk.cell(r, c).value == inv.cell(r, c).value, (r, c)
    for p in plan:
        assert float(chk.cell(p["rij"], 13).value) == p["nieuw"], p
    shutil.move(tmp, args.bestand)
    print(f"\ngeschreven: {args.bestand}\nback-up:    {backup}")
    print("vergeet niet: het team werkt in Google Sheets — upload het bestand of neem de prijzen daar over.")


if __name__ == "__main__":
    main()
